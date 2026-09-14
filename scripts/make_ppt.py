#!/usr/bin/env python3
"""Generate a 20-slide Review-2 presentation (STAG-STI FSCIL) as a .pptx.

Pulls real numbers from checkpoints/<dataset>/incremental_results.json, renders
a couple of charts with matplotlib, and lays out a styled deck. Run:

    .venv/bin/python scripts/make_ppt.py
"""
import json, os
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
ASSET = os.path.join(ROOT, "results", "ppt_assets")
os.makedirs(ASSET, exist_ok=True)

# ---------------- palette ----------------
INK      = RGBColor(0x14, 0x1B, 0x2E)   # near-black navy
NAVY     = RGBColor(0x1E, 0x3A, 0x5F)   # deep blue
BLUE     = RGBColor(0x2D, 0x6C, 0xDF)   # accent blue
TEAL     = RGBColor(0x17, 0x9E, 0x8A)   # accent teal
AMBER    = RGBColor(0xE0, 0x8A, 0x1E)   # accent amber
RED      = RGBColor(0xC0, 0x39, 0x2B)
SLATE    = RGBColor(0x5A, 0x67, 0x78)   # secondary text
LIGHT    = RGBColor(0xF4, 0xF6, 0xFA)   # panel bg
CARD     = RGBColor(0xEB, 0xEF, 0xF6)
WHITE    = RGBColor(0xFF, 0xFF, 0xFF)
LINE     = RGBColor(0xD3, 0xDA, 0xE6)

FONT = "Calibri"
FONTH = "Calibri"

# ---------------- load data ----------------
def load(d):
    return json.load(open(os.path.join(ROOT, "checkpoints", d, "incremental_results.json")))

DATA = {d: load(d) for d in ["cifar100", "miniimagenet", "cub200"]}
NAMES = {"cifar100": "CIFAR-100", "miniimagenet": "miniImageNet", "cub200": "CUB-200"}
COLORS = {"cifar100": "#2D6CDF", "miniimagenet": "#179E8A", "cub200": "#E08A1E"}

# ================= chart rendering =================
def chart_accuracy_curves():
    fig, ax = plt.subplots(figsize=(7.6, 4.2), dpi=200)
    for d in ["cifar100", "miniimagenet", "cub200"]:
        res = DATA[d]["results"]
        xs = [r["session"] for r in res]
        ys = [r["accuracy"] * 100 for r in res]
        ax.plot(xs, ys, marker="o", lw=2.4, ms=5, color=COLORS[d], label=NAMES[d])
    ax.set_xlabel("Incremental session", fontsize=11)
    ax.set_ylabel("Overall accuracy (%)", fontsize=11)
    ax.set_title("Accuracy across incremental sessions", fontsize=13, weight="bold", color="#141B2E")
    ax.grid(True, alpha=0.25)
    ax.legend(frameon=False, fontsize=10)
    ax.set_ylim(35, 82)
    for s in ["top", "right"]:
        ax.spines[s].set_visible(False)
    fig.tight_layout()
    p = os.path.join(ASSET, "curves.png")
    fig.savefig(p, bbox_inches="tight"); plt.close(fig)
    return p

def chart_base_vs_novel():
    fig, ax = plt.subplots(figsize=(7.6, 4.2), dpi=200)
    labels = [NAMES[d] for d in ["cifar100", "miniimagenet", "cub200"]]
    ab = [DATA[d]["A_B"] for d in ["cifar100", "miniimagenet", "cub200"]]
    an = [DATA[d]["A_N"] for d in ["cifar100", "miniimagenet", "cub200"]]
    x = range(len(labels)); w = 0.36
    ax.bar([i - w/2 for i in x], ab, w, label="A_B (base-only)", color="#2D6CDF")
    ax.bar([i + w/2 for i in x], an, w, label="A_N (novel-only)", color="#C0392B")
    for i, v in enumerate(ab):
        ax.text(i - w/2, v + 1, f"{v:.1f}", ha="center", fontsize=9, weight="bold")
    for i, v in enumerate(an):
        ax.text(i + w/2, v + 1, f"{v:.1f}", ha="center", fontsize=9, weight="bold")
    ax.set_xticks(list(x)); ax.set_xticklabels(labels, fontsize=11)
    ax.set_ylabel("Final-session accuracy (%)", fontsize=11)
    ax.set_title("The novel-class gap (A_B vs A_N)", fontsize=13, weight="bold", color="#141B2E")
    ax.legend(frameon=False, fontsize=10); ax.set_ylim(0, 80)
    for s in ["top", "right"]:
        ax.spines[s].set_visible(False)
    ax.grid(True, axis="y", alpha=0.25)
    fig.tight_layout()
    p = os.path.join(ASSET, "base_novel.png")
    fig.savefig(p, bbox_inches="tight"); plt.close(fig)
    return p

def chart_improvement():
    fig, ax = plt.subplots(figsize=(7.6, 4.2), dpi=200)
    labels = ["CIFAR-100", "miniImageNet", "CUB-200"]
    base = [49.53, 40.25, 49.71]
    best = [51.13, 41.86, 50.91]
    x = range(len(labels)); w = 0.36
    ax.bar([i - w/2 for i in x], base, w, label="Baseline (inductive)", color="#5A6778")
    ax.bar([i + w/2 for i in x], best, w, label="Best eval-only config", color="#179E8A")
    for i in x:
        ax.text(i + w/2, best[i] + 0.5, f"+{best[i]-base[i]:.1f}", ha="center", fontsize=9,
                weight="bold", color="#179E8A")
    ax.set_xticks(list(x)); ax.set_xticklabels(labels, fontsize=11)
    ax.set_ylabel("Final accuracy (%)", fontsize=11)
    ax.set_title("Eval-only improvement (no architecture change)", fontsize=13, weight="bold", color="#141B2E")
    ax.legend(frameon=False, fontsize=10); ax.set_ylim(0, 62)
    for s in ["top", "right"]:
        ax.spines[s].set_visible(False)
    ax.grid(True, axis="y", alpha=0.25)
    fig.tight_layout()
    p = os.path.join(ASSET, "improve.png")
    fig.savefig(p, bbox_inches="tight"); plt.close(fig)
    return p

CURVES = chart_accuracy_curves()
BASENOVEL = chart_base_vs_novel()
IMPROVE = chart_improvement()

# ================= slide helpers =================
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]

def slide():
    return prs.slides.add_slide(BLANK)

def rect(s, x, y, w, h, fill, line=None, line_w=None):
    from pptx.enum.shapes import MSO_SHAPE
    shp = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    shp.fill.solid(); shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line; shp.line.width = line_w or Pt(1)
    shp.shadow.inherit = False
    return shp

def txt(s, x, y, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
        space_after=4, line_spacing=1.0):
    """runs: list of paragraphs; each paragraph = list of (text, size, color, bold, font)."""
    tb = s.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame; tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = Pt(2); tf.margin_right = Pt(2)
    tf.margin_top = Pt(1); tf.margin_bottom = Pt(1)
    for i, para in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align; p.space_after = Pt(space_after); p.space_before = Pt(0)
        p.line_spacing = line_spacing
        for (t, sz, col, bold, fnt) in para:
            r = p.add_run(); r.text = t
            r.font.size = Pt(sz); r.font.color.rgb = col
            r.font.bold = bold; r.font.name = fnt
    return tb

def P(text, size, color=INK, bold=False, font=FONT):
    return [(text, size, color, bold, font)]

def bullets(s, x, y, w, h, items, size=15, gap=8, color=INK, marker="●",
            marker_col=BLUE, line_spacing=1.02):
    tb = s.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame; tf.word_wrap = True
    tf.margin_left = Pt(2); tf.margin_right = Pt(2)
    for i, it in enumerate(items):
        if isinstance(it, tuple):
            txt_str, lvl = it
        else:
            txt_str, lvl = it, 0
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(gap); p.space_before = Pt(0); p.line_spacing = line_spacing
        if lvl == 0:
            r = p.add_run(); r.text = marker + "  "
            r.font.size = Pt(size); r.font.color.rgb = marker_col; r.font.name = FONT; r.font.bold = True
        else:
            r = p.add_run(); r.text = "      –  "
            r.font.size = Pt(size - 1); r.font.color.rgb = SLATE; r.font.name = FONT
        # allow inline bold via **
        parts = txt_str.split("**")
        for j, seg in enumerate(parts):
            if not seg:
                continue
            r = p.add_run(); r.text = seg
            r.font.size = Pt(size - (1 if lvl else 0))
            r.font.color.rgb = color; r.font.name = FONT
            r.font.bold = (j % 2 == 1)
    return tb

def header(s, kicker, title, idx):
    rect(s, 0, 0, SW, Inches(1.15), WHITE)
    rect(s, 0, 0, Inches(0.18), Inches(1.15), BLUE)
    txt(s, Inches(0.55), Inches(0.16), Inches(10.5), Inches(0.34),
        [P(kicker.upper(), 12, BLUE, True)])
    txt(s, Inches(0.55), Inches(0.44), Inches(11.6), Inches(0.62),
        [P(title, 27, INK, True)])
    # slide number chip
    rect(s, SW - Inches(0.95), Inches(0.30), Inches(0.55), Inches(0.55), CARD)
    txt(s, SW - Inches(0.95), Inches(0.30), Inches(0.55), Inches(0.55),
        [P(f"{idx:02d}", 15, NAVY, True)], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    rect(s, 0, Inches(1.15), SW, Pt(2), LINE)

def footer(s):
    txt(s, Inches(0.55), SH - Inches(0.42), Inches(9), Inches(0.3),
        [P("STAG-STI · Few-Shot Class-Incremental Learning · Review 2", 9, SLATE, False)])

def table(s, x, y, w, headers, rows, col_w=None, row_h=Inches(0.34),
          head_fill=NAVY, fs=12, hfs=12, zebra=True,
          bold_first=False, highlight_rows=None):
    highlight_rows = highlight_rows or {}
    n = len(headers)
    if col_w is None:
        col_w = [w // n] * n
    # header
    cx = x
    hh = Inches(0.44)
    for j, htext in enumerate(headers):
        c = rect(s, cx, y, col_w[j], hh, head_fill)
        txt(s, cx, y, col_w[j], hh, [P(htext, hfs, WHITE, True)],
            align=PP_ALIGN.CENTER if j else PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE)
        cx += col_w[j]
    ry = y + hh
    for i, row in enumerate(rows):
        cx = x
        base = WHITE if i % 2 == 0 else LIGHT
        if i in highlight_rows:
            base = highlight_rows[i]
        for j, cell in enumerate(row):
            rect(s, cx, ry, col_w[j], row_h, base, line=LINE, line_w=Pt(0.5))
            bold = (bold_first and j == 0) or (i in highlight_rows)
            col = INK
            txt(s, cx + Inches(0.03), ry, col_w[j] - Inches(0.06), row_h,
                [P(str(cell), fs, col, bold)],
                align=PP_ALIGN.CENTER if j else PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE)
            cx += col_w[j]
        ry += row_h
    return ry

def statcard(s, x, y, w, h, big, label, accent=BLUE, sub=None):
    rect(s, x, y, w, h, LIGHT, line=LINE, line_w=Pt(0.75))
    rect(s, x, y, w, Inches(0.09), accent)
    txt(s, x, y + Inches(0.18), w, Inches(0.6),
        [P(big, 30, accent, True)], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    txt(s, x, y + h - Inches(0.78), w, Inches(0.7),
        [P(label, 12, INK, True)] if not sub else [P(label, 12, INK, True)],
        align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.TOP)
    if sub:
        txt(s, x + Inches(0.05), y + h - Inches(0.42), w - Inches(0.1), Inches(0.36),
            [P(sub, 9, SLATE, False)], align=PP_ALIGN.CENTER)

# ==================================================================
# SLIDE 1 — TITLE
# ==================================================================
s = slide()
rect(s, 0, 0, SW, SH, INK)
rect(s, 0, 0, SW, Inches(0.22), BLUE)
rect(s, 0, SH - Inches(0.22), SW, Inches(0.22), TEAL)
txt(s, Inches(0.9), Inches(1.5), Inches(11.5), Inches(0.5),
    [P("FEW-SHOT CLASS-INCREMENTAL LEARNING", 15, TEAL, True)])
txt(s, Inches(0.9), Inches(2.05), Inches(11.6), Inches(1.7),
    [P("STAG-STI", 54, WHITE, True)])
txt(s, Inches(0.9), Inches(3.05), Inches(11.6), Inches(1.4),
    [[("Spatio-Temporal Graph Integration with Contrastive Pre-Conditioning", 24, WHITE, True, FONT)],
     [("for Forward-Compatible Few-Shot Class-Incremental Learning", 24, WHITE, True, FONT)]],
    line_spacing=1.05)
rect(s, Inches(0.92), Inches(4.35), Inches(4.4), Pt(2), TEAL)
txt(s, Inches(0.9), Inches(4.55), Inches(11.6), Inches(1.2),
    [[("Evaluated on CIFAR-100 · miniImageNet · CUB-200 + Cross-Domain Transfer", 16, RGBColor(0xC9,0xD4,0xE4), False, FONT)],
     [("Official CEC / FACT splits · full ablation & improvement study", 14, SLATE, False, FONT)]],
    line_spacing=1.3)
txt(s, Inches(0.9), Inches(6.2), Inches(11.6), Inches(0.7),
    [[("Review 2  ·  Project Presentation", 15, WHITE, True, FONT)]])

# ==================================================================
# SLIDE 2 — PROBLEM / MOTIVATION
# ==================================================================
s = slide(); header(s, "Motivation", "The Problem: Learning Continuously from Few Examples", 2); footer(s)
bullets(s, Inches(0.55), Inches(1.45), Inches(6.5), Inches(5),
    [
     "**Real-world models must keep learning** new categories after deployment — a medical system meets a new disease, a retail model a new product.",
     "New classes arrive with **only a handful of labelled examples** (5-shot), not the thousands used in base training.",
     "**Catastrophic forgetting:** naively fine-tuning on new classes destroys knowledge of old ones.",
     "**Overfitting:** 5 examples are far too few to train a deep network from scratch.",
     "FSCIL must balance **stability** (retain old) and **plasticity** (absorb new) at once.",
    ], size=15, gap=13)
# right panel
rect(s, Inches(7.35), Inches(1.5), Inches(5.35), Inches(4.7), LIGHT, line=LINE, line_w=Pt(1))
txt(s, Inches(7.6), Inches(1.7), Inches(4.9), Inches(0.5), [P("The FSCIL dilemma", 15, NAVY, True)])
txt(s, Inches(7.6), Inches(2.35), Inches(4.9), Inches(1.2),
    [[("Stability  ⇄  Plasticity", 20, INK, True, FONT)]], align=PP_ALIGN.CENTER)
rect(s, Inches(7.6), Inches(3.15), Inches(4.85), Pt(1.5), LINE)
bullets(s, Inches(7.6), Inches(3.35), Inches(4.85), Inches(2.7),
    [
     "**Too much plasticity** → forgets base classes",
     "**Too much stability** → can't learn novel classes",
     "**Few shots** → prototypes are noisy & biased",
     "**No replay** → cannot revisit old data",
    ], size=13, gap=11, marker="▸", marker_col=AMBER)

# ==================================================================
# SLIDE 3 — OBJECTIVE & SCOPE
# ==================================================================
s = slide(); header(s, "Objective", "Research Objective & Project Scope", 3); footer(s)
txt(s, Inches(0.55), Inches(1.4), Inches(12), Inches(0.8),
    [[("Goal: ", 17, TEAL, True, FONT),
      ("Design and validate a forward-compatible FSCIL system that maximises novel-class learning without forgetting — and quantify every design choice.", 17, INK, False, FONT)]],
    line_spacing=1.1)
# three scope cards
cards = [
    ("Architecture", "Implement the full STAG-STI 8-stage pipeline: contrastive pre-conditioning + topology-aware graph attention + age-decaying memory.", BLUE),
    ("Evaluation", "Benchmark on 3 standard datasets on official splits, plus a cross-domain transfer test. Report the full FSCIL metric set.", TEAL),
    ("Analysis", "Ablate every component, run an evidence-driven improvement study, and identify the true accuracy bottleneck.", AMBER),
]
cw = Inches(3.95); gap = Inches(0.28); x0 = Inches(0.55); y0 = Inches(2.45); ch = Inches(2.5)
for i, (t, d, c) in enumerate(cards):
    x = x0 + i * (cw + gap)
    rect(s, x, y0, cw, ch, WHITE, line=LINE, line_w=Pt(1))
    rect(s, x, y0, cw, Inches(0.5), c)
    txt(s, x, y0, cw, Inches(0.5), [P(t, 15, WHITE, True)], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    txt(s, x + Inches(0.2), y0 + Inches(0.65), cw - Inches(0.4), Inches(1.8),
        [P(d, 13, INK, False)], line_spacing=1.12)
txt(s, Inches(0.55), Inches(5.4), Inches(12.2), Inches(1.5),
    [[("Key design decisions:  ", 13, NAVY, True, FONT),
      ("official CEC/FACT splits · exact official few-shot support samples · deterministic fantasy views · CUB novel sessions 10-way · frozen backbone at incremental time · no exemplar replay.", 13, SLATE, False, FONT)]],
    line_spacing=1.15)

# ==================================================================
# SLIDE 4 — BACKGROUND: FSCIL PROTOCOL
# ==================================================================
s = slide(); header(s, "Background", "FSCIL Protocol: How Evaluation Works", 4); footer(s)
bullets(s, Inches(0.55), Inches(1.45), Inches(6.4), Inches(5),
    [
     "**Session 0 (base):** many classes, abundant data — trains the backbone.",
     "**Sessions 1…T (incremental):** each adds N new classes with K shots (N-way K-shot).",
     "**Disjoint label spaces:** a class appears in exactly one session.",
     "**Global evaluation:** after each session, test over *all* classes seen so far.",
     "**Frozen backbone:** after base training, weights never change — only prototypes are added.",
     "**Nearest-prototype classification:** cosine similarity to stored class prototypes.",
    ], size=14.5, gap=11)
rect(s, Inches(7.25), Inches(1.5), Inches(5.5), Inches(4.7), LIGHT, line=LINE, line_w=Pt(1))
txt(s, Inches(7.5), Inches(1.68), Inches(5), Inches(0.4), [P("Session timeline", 14, NAVY, True)])
# mini timeline
tl_y = Inches(2.5)
labels = ["S0\nbase", "S1", "S2", "…", "S_T"]
tlw = Inches(0.85); tgap = Inches(0.18); tx = Inches(7.55)
cols = [NAVY, BLUE, BLUE, SLATE, TEAL]
for i, l in enumerate(labels):
    x = tx + i * (tlw + tgap)
    rect(s, x, tl_y, tlw, Inches(0.85), cols[i])
    txt(s, x, tl_y, tlw, Inches(0.85), [P(l, 12, WHITE, True)], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
txt(s, Inches(7.5), Inches(3.7), Inches(5), Inches(0.4),
    [P("classes seen grows each session →", 11, SLATE, True)])
bullets(s, Inches(7.5), Inches(4.2), Inches(5), Inches(1.9),
    [
     "Accuracy naturally **declines** as more classes compete.",
     "Success = a **slow, graceful** decline, not a cliff.",
    ], size=13, gap=10, marker="▸", marker_col=AMBER)

# ==================================================================
# SLIDE 5 — ARCHITECTURE OVERVIEW
# ==================================================================
s = slide(); header(s, "Architecture", "STAG-STI Pipeline — End to End", 5); footer(s)
txt(s, Inches(0.55), Inches(1.35), Inches(12), Inches(0.5),
    [P("Eight stages turn a few raw images into a robust, forward-compatible class prototype:", 14, SLATE, False)])
stages = [
    ("1  Fantasy Views ×M", "Deterministic multi-view augmentation", BLUE),
    ("2  Frozen Backbone fθ", "ResNet-18 → 512-d features", NAVY),
    ("3  Projection Wₚ", "512-d → 256-d bottleneck", BLUE),
    ("4  SupCon Head", "Contrastive pre-conditioning", TEAL),
    ("5  Topology Scorer gφ", "Learns the view/class graph", TEAL),
    ("6  GATv2 Attention", "Message passing over the graph", BLUE),
    ("7  Prototype Pooling", "Aggregate → one prototype", NAVY),
    ("8  STI Memory Hₜ", "Age-decaying stability prior", AMBER),
]
cw = Inches(2.95); ch = Inches(1.15); gx = Inches(0.22); gy = Inches(0.28)
x0 = Inches(0.55); y0 = Inches(2.0)
for i, (t, d, c) in enumerate(stages):
    r, col = divmod(i, 4)
    x = x0 + col * (cw + gx); y = y0 + r * (ch + gy)
    rect(s, x, y, cw, ch, WHITE, line=LINE, line_w=Pt(1))
    rect(s, x, y, Inches(0.1), ch, c)
    txt(s, x + Inches(0.22), y + Inches(0.14), cw - Inches(0.3), Inches(0.5), [P(t, 13.5, INK, True)])
    txt(s, x + Inches(0.22), y + Inches(0.58), cw - Inches(0.3), Inches(0.5), [P(d, 11, SLATE, False)], line_spacing=1.0)
    if col < 3:
        txt(s, x + cw - Inches(0.02), y + Inches(0.35), Inches(0.26), Inches(0.5),
            [P("→", 18, SLATE, True)], align=PP_ALIGN.CENTER)
txt(s, Inches(0.55), Inches(6.35), Inches(12), Inches(0.6),
    [[("Final stage: ", 13, NAVY, True, FONT),
      ("a cosine classifier scores queries against all prototypes. Queries skip the graph — only prototypes are graph-refined, keeping inference cheap.", 13, SLATE, False, FONT)]],
    line_spacing=1.1)

# ==================================================================
# SLIDE 6 — COMPONENT 1: CONTRASTIVE
# ==================================================================
s = slide(); header(s, "Component 1 of 3", "Contrastive Pre-Conditioning (SupCon)", 6); footer(s)
bullets(s, Inches(0.55), Inches(1.5), Inches(6.5), Inches(5),
    [
     "A **supervised contrastive** (SupCon) head shapes the feature space during base training.",
     "Pulls same-class features together, pushes different classes apart on a hypersphere.",
     "**Why it matters for FSCIL:** a well-spread, well-clustered space means a 5-shot prototype lands in a meaningful region.",
     "Operates on the raw 512-d backbone features (in_dim = backbone_out_dim).",
     "Temperature 0.15, weight λ_supcon = 0.3 — tuned to avoid over-tight clustering that harms novel classes.",
    ], size=14.5, gap=12)
rect(s, Inches(7.35), Inches(1.55), Inches(5.35), Inches(4.4), LIGHT, line=LINE, line_w=Pt(1))
txt(s, Inches(7.6), Inches(1.75), Inches(5), Inches(0.4), [P("Intuition", 14, TEAL, True)])
txt(s, Inches(7.6), Inches(2.35), Inches(4.85), Inches(2),
    [[("Better-organised base features", 16, INK, True, FONT)],
     [("↓", 20, TEAL, True, FONT)],
     [("Novel prototypes generalise better", 16, INK, True, FONT)],
     [("↓", 20, TEAL, True, FONT)],
     [("Higher novel-class accuracy", 16, TEAL, True, FONT)]],
    align=PP_ALIGN.CENTER, line_spacing=1.05, space_after=6)

# ==================================================================
# SLIDE 7 — COMPONENT 2: TOPOLOGY / GAT
# ==================================================================
s = slide(); header(s, "Component 2 of 3", "Topology-Aware Graph Attention (GATv2)", 7); footer(s)
bullets(s, Inches(0.55), Inches(1.5), Inches(6.5), Inches(5),
    [
     "A **topology scorer gφ** predicts how strongly nodes (views/classes) should be connected.",
     "**GATv2** performs attention-weighted message passing over that learned graph.",
     "Multiple fantasy views of the same class exchange information → a **denoised, consensus prototype**.",
     "Supervised by a graph BCE loss (λ_graph = 0.25) so the learned adjacency is meaningful, not arbitrary.",
     "Dropout inside the GAT fights Phase-B overfitting on the small episodic batches.",
    ], size=14.5, gap=12)
rect(s, Inches(7.35), Inches(1.55), Inches(5.35), Inches(4.4), LIGHT, line=LINE, line_w=Pt(1))
txt(s, Inches(7.6), Inches(1.75), Inches(5), Inches(0.4), [P("What the graph gives", 14, BLUE, True)])
bullets(s, Inches(7.6), Inches(2.35), Inches(4.85), Inches(3.4),
    [
     "**Noise averaging** across views of one class",
     "**Structure awareness** — related classes inform each other",
     "**Query-independent** — only prototypes pass through the graph",
     "Isolated in the ablation study (see Slide 15)",
    ], size=13.5, gap=13, marker="▸", marker_col=BLUE)

# ==================================================================
# SLIDE 8 — COMPONENT 3: STI MEMORY
# ==================================================================
s = slide(); header(s, "Component 3 of 3", "Spatio-Temporal Memory (Age-Decaying Prior)", 8); footer(s)
bullets(s, Inches(0.55), Inches(1.5), Inches(6.5), Inches(5),
    [
     "The **STI memory Hₜ** stores every class prototype and controls how much each can still change.",
     "An **age-based stability prior**: older classes harden (β → 1), new classes stay plastic (β ≈ 0).",
     "Controlled by a gate Γ and reset gate R_gate; hardening speed κ = 2.0, offset κ₀ = 3.0.",
     "**Directly targets the stability–plasticity trade-off:** protect the old, adapt the new.",
     "Enables incremental additions with **no replay and no backbone updates**.",
    ], size=14.5, gap=12)
rect(s, Inches(7.35), Inches(1.55), Inches(5.35), Inches(4.4), LIGHT, line=LINE, line_w=Pt(1))
txt(s, Inches(7.6), Inches(1.75), Inches(5), Inches(0.4), [P("Age → stability", 14, AMBER, True)])
# simple decay illustration
txt(s, Inches(7.6), Inches(2.4), Inches(4.85), Inches(3.2),
    [[("New class", 14, INK, True, FONT), ("   β ≈ 0  (plastic)", 14, TEAL, False, FONT)],
     [("     ↓  time / sessions", 12, SLATE, False, FONT)],
     [("Aged class", 14, INK, True, FONT), ("   β → 1  (frozen)", 14, RED, False, FONT)],
     [("", 6, INK, False, FONT)],
     [("Result: recent classes adapt,", 13, INK, False, FONT)],
     [("established classes are protected.", 13, INK, False, FONT)]],
    line_spacing=1.15, space_after=8)

# ==================================================================
# SLIDE 9 — DATASETS & PROTOCOL
# ==================================================================
s = slide(); header(s, "Experimental Setup", "Datasets & Evaluation Protocol", 9); footer(s)
headers = ["Dataset", "Resolution", "Base / Novel", "Sessions", "Backbone"]
rows = [
    ["CIFAR-100", "32×32", "60 / 40", "8 × (5-way 5-shot)", "ResNet-18 (CIFAR stem)"],
    ["miniImageNet", "84×84", "60 / 40", "8 × (5-way 5-shot)", "ResNet-18"],
    ["CUB-200-2011", "224×224", "100 / 100", "10 × (10-way 5-shot)", "ResNet-18 (ImageNet-pretrained)"],
    ["Cross-domain", "84×84", "mini → CUB", "10 × (10-way 5-shot)", "miniImageNet backbone (frozen)"],
]
cw = [Inches(2.5), Inches(1.9), Inches(2.0), Inches(3.0), Inches(3.0)]
table(s, Inches(0.55), Inches(1.5), Inches(12.2), headers, rows, col_w=cw,
      row_h=Inches(0.52), fs=12.5, hfs=12.5,
      highlight_rows={3: CARD})
bullets(s, Inches(0.55), Inches(4.55), Inches(12), Inches(2.2),
    [
     "**Official CEC / FACT splits** used throughout — vendored under fscil/splits/ for exact reproducibility.",
     "Incremental support uses the **exact official few-shot samples**, not random draws.",
     "Cross-domain is this project's own protocol: base representation from miniImageNet, novel classes from CUB.",
    ], size=13.5, gap=10, marker="●", marker_col=TEAL)

# ==================================================================
# SLIDE 10 — METRICS
# ==================================================================
s = slide(); header(s, "Experimental Setup", "Evaluation Metrics", 10); footer(s)
metrics = [
    ("Aᵢ", "Accuracy after session i over all classes seen", BLUE),
    ("A₀ / A_T", "Base (first) and final-session accuracy", NAVY),
    ("PD ↓", "Performance Drop = A₀ − A_T (forgetting)", RED),
    ("A_B / A_N", "Base-only / novel-only accuracy at the end", AMBER),
    ("HM", "Harmonic mean of A_B and A_N (balance)", TEAL),
    ("Avg", "Mean accuracy across all sessions", SLATE),
]
cw = Inches(3.95); ch = Inches(1.55); gx = Inches(0.28); gy = Inches(0.3)
x0 = Inches(0.55); y0 = Inches(1.55)
for i, (sym, desc, c) in enumerate(metrics):
    r, col = divmod(i, 3)
    x = x0 + col * (cw + gx); y = y0 + r * (ch + gy)
    rect(s, x, y, cw, ch, WHITE, line=LINE, line_w=Pt(1))
    rect(s, x, y, cw, Inches(0.08), c)
    txt(s, x + Inches(0.2), y + Inches(0.2), cw - Inches(0.4), Inches(0.55),
        [P(sym, 22, c, True)])
    txt(s, x + Inches(0.2), y + Inches(0.78), cw - Inches(0.4), Inches(0.7),
        [P(desc, 12.5, INK, False)], line_spacing=1.05)
txt(s, Inches(0.55), Inches(5.7), Inches(12), Inches(0.9),
    [[("Why HM matters most: ", 14, TEAL, True, FONT),
      ("a model can inflate overall accuracy by acquiescing to base classes. HM only rises when BOTH base and novel accuracy are healthy — it is the honest FSCIL score.", 14, INK, False, FONT)]],
    line_spacing=1.12)

# ==================================================================
# SLIDE 11 — IMPLEMENTATION
# ==================================================================
s = slide(); header(s, "Experimental Setup", "Implementation & Training Recipe", 11); footer(s)
bullets(s, Inches(0.55), Inches(1.5), Inches(6.5), Inches(5),
    [
     "**Framework:** PyTorch, Apple-Silicon (MPS) backend.",
     "**Phase A — backbone:** cross-entropy pretraining, 30 epochs, cosine LR, label smoothing 0.1, Mixup/CutMix, random erasing, early stopping.",
     "**Phase B — STAG-STI:** 50 episodic epochs, 100 episodes/epoch, 15-way 5-shot 5-query episodes, composite loss (cls + SupCon + graph + stability).",
     "**Phase C — incremental eval:** freeze everything; add prototypes session by session.",
     "**Reproducibility:** fixed seed 42, deterministic fantasy views, namespaced checkpoints.",
    ], size=13.5, gap=11)
rect(s, Inches(7.35), Inches(1.55), Inches(5.35), Inches(4.5), LIGHT, line=LINE, line_w=Pt(1))
txt(s, Inches(7.6), Inches(1.75), Inches(5), Inches(0.4), [P("Composite Phase-B loss", 14, NAVY, True)])
bullets(s, Inches(7.6), Inches(2.4), Inches(4.9), Inches(3.5),
    [
     "**L_cls** — query classification (λ = 1.0)",
     "**L_supcon** — contrastive (λ = 0.3)",
     "**L_graph** — topology BCE (λ = 0.25)",
     "**L_stab** — memory + gate reg. (λ = 0.1)",
    ], size=13.5, gap=13, marker="▸", marker_col=NAVY)
txt(s, Inches(7.6), Inches(5.15), Inches(4.9), Inches(0.8),
    [P("Weights tuned on a held-out val split to balance clustering vs novel-class flexibility.", 11.5, SLATE, False)], line_spacing=1.1)

# ==================================================================
# SLIDE 12 — RESULTS: IN-DOMAIN (main table + curve)
# ==================================================================
s = slide(); header(s, "Results", "In-Domain Performance (Official Splits)", 12); footer(s)
headers = ["Dataset", "A₀", "A_T", "PD↓", "Avg", "A_B", "A_N", "HM"]
rows = []
for d in ["cifar100", "miniimagenet", "cub200"]:
    r = DATA[d]
    a0 = r["results"][0]["accuracy"] * 100
    at = r["results"][-1]["accuracy"] * 100
    rows.append([NAMES[d], f"{a0:.1f}", f"{at:.1f}", f"{r['PD']:.1f}",
                 f"{r['avg_accuracy']:.1f}", f"{r['A_B']:.1f}", f"{r['A_N']:.1f}", f"{r['harmonic_mean']:.1f}"])
cw = [Inches(1.9)] + [Inches(0.72)] * 7
table(s, Inches(0.55), Inches(1.5), Inches(7.05), headers, rows, col_w=cw,
      row_h=Inches(0.5), fs=12.5, hfs=12, bold_first=True)
s.shapes.add_picture(CURVES, Inches(0.55), Inches(3.55), width=Inches(6.9))
# right notes
rect(s, Inches(7.75), Inches(1.5), Inches(5.0), Inches(5.0), LIGHT, line=LINE, line_w=Pt(1))
txt(s, Inches(7.95), Inches(1.7), Inches(4.6), Inches(0.4), [P("Reading the results", 14, NAVY, True)])
bullets(s, Inches(7.95), Inches(2.3), Inches(4.6), Inches(4),
    [
     "Base accuracy is **strong** (66–78%).",
     "Decline is **graceful**, not a cliff — the memory prior is working.",
     "CUB holds up best (HM **42.6**) thanks to its ImageNet-pretrained backbone.",
     "miniImageNet is hardest (HM 25.4) — weakest backbone, hardest novel transfer.",
     "Numbers are **single-run, first-pass** on official splits — internally consistent, not seed-averaged.",
    ], size=13, gap=11, marker="▸", marker_col=BLUE)

# ==================================================================
# SLIDE 13 — TRANSDUCTIVE RECTIFICATION
# ==================================================================
s = slide(); header(s, "Results", "Improvement: Transductive Prototype Rectification", 13); footer(s)
txt(s, Inches(0.55), Inches(1.35), Inches(12), Inches(0.55),
    [P("A BD-CSPN-style transductive step refines novel prototypes using unlabelled test features — the only lever that improved every metric.", 13.5, SLATE, False)],
    line_spacing=1.1)
headers = ["Dataset", "Final (ind.)", "Final (transd.)", "A_N (ind.)", "A_N (transd.)", "HM (ind.)", "HM (transd.)"]
rows = [
    ["CIFAR-100", "49.53", "50.17", "20.75", "21.45", "31.87", "32.76"],
    ["miniImageNet", "40.25", "40.89", "16.45", "16.53", "25.44", "25.64"],
    ["CUB-200", "49.71", "50.91", "30.85", "33.52", "42.64", "45.06"],
]
cw = [Inches(1.9)] + [Inches(1.72)] * 6
table(s, Inches(0.55), Inches(2.1), Inches(12.2), headers, rows, col_w=cw,
      row_h=Inches(0.52), fs=12.5, hfs=11.5, bold_first=True,
      highlight_rows={2: CARD})
bullets(s, Inches(0.55), Inches(4.5), Inches(12), Inches(2.2),
    [
     "**Every metric improves** on all three datasets — no base↔novel trade.",
     "**Biggest gain on CUB:** A_N +2.7, HM +2.4 (transductive refinement helps most when novel prototypes are noisiest).",
     "Combined with **full-data base prototypes** on weak backbones, best-config gains reach **+1.2 to +1.6 overall** (next slide).",
    ], size=13.5, gap=11, marker="●", marker_col=TEAL)

# ==================================================================
# SLIDE 14 — CROSS-DOMAIN + IMPROVEMENT SUMMARY
# ==================================================================
s = slide(); header(s, "Results", "Cross-Domain Transfer & Best-Config Gains", 14); footer(s)
# left: cross domain stat cards
txt(s, Inches(0.55), Inches(1.35), Inches(6), Inches(0.4), [P("Cross-domain: miniImageNet → CUB", 15, NAVY, True)])
statcard(s, Inches(0.55), Inches(1.85), Inches(2.9), Inches(1.7), "66.7%", "Base (miniImageNet)", BLUE)
statcard(s, Inches(3.6), Inches(1.85), Inches(2.9), Inches(1.7), "44.3%", "Final (all seen)", AMBER)
statcard(s, Inches(0.55), Inches(3.75), Inches(2.9), Inches(1.7), "22.4", "PD (forgetting)", RED)
statcard(s, Inches(3.6), Inches(3.75), Inches(2.9), Inches(1.7), "14.8", "HM", TEAL)
txt(s, Inches(0.55), Inches(5.6), Inches(6), Inches(1),
    [P("Domain shift is hard: A_N drops to 8.4 — features tuned on natural images transfer only partially to fine-grained birds.", 12.5, SLATE, False)], line_spacing=1.12)
# right: improvement chart
s.shapes.add_picture(IMPROVE, Inches(6.9), Inches(1.5), width=Inches(6.0))
txt(s, Inches(6.9), Inches(5.7), Inches(6), Inches(1),
    [[("Best eval-only stack: ", 12.5, TEAL, True, FONT),
      ("full-data base prototypes (where the backbone is weak) + transductive rectification — no architecture change, no retraining.", 12.5, INK, False, FONT)]],
    line_spacing=1.12)

# ==================================================================
# SLIDE 15 — ABLATION STUDY
# ==================================================================
s = slide(); header(s, "Analysis", "Ablation Study (CIFAR-100)", 15); footer(s)
txt(s, Inches(0.55), Inches(1.35), Inches(12), Inches(0.5),
    [P("Each component removed in isolation, sharing the same frozen backbone. HM is the deciding metric.", 13.5, SLATE, False)])
headers = ["Configuration", "A₀", "A_T", "A_B", "A_N", "HM"]
rows = [
    ["Full model", "78.22", "49.53", "68.72", "20.75", "31.87"],
    ["− contrastive (SupCon)", "78.53", "49.28", "69.97", "18.25", "28.95"],
    ["− topology (GAT graph)", "77.75", "49.69", "70.98", "17.75", "28.40"],
    ["− age-decay (STI)", "78.50", "49.46", "70.15", "18.43", "29.18"],
]
cw = [Inches(4.4)] + [Inches(1.45)] * 5
table(s, Inches(0.55), Inches(2.05), Inches(11.65), headers, rows, col_w=cw,
      row_h=Inches(0.52), fs=13, hfs=12.5, bold_first=True,
      highlight_rows={0: CARD})
bullets(s, Inches(0.55), Inches(4.6), Inches(12), Inches(2.2),
    [
     "**Removing any component lowers HM** — every piece earns its place.",
     "Ablations **raise A_B but crash A_N**: without the full model, the system defaults to over-favouring base classes.",
     "The components' shared job is to **protect novel-class accuracy** — exactly the FSCIL bottleneck.",
    ], size=13.5, gap=11, marker="●", marker_col=BLUE)

# ==================================================================
# SLIDE 16 — CLOSER STUDY
# ==================================================================
s = slide(); header(s, "Analysis", "CLOSER-Style Base Objective (Transfer Study)", 16); footer(s)
txt(s, Inches(0.55), Inches(1.35), Inches(12), Inches(0.5),
    [P("Testing an alternative base objective (Oh et al., ECCV'24) that spreads features and reduces inter-class distance for reusable low/mid features.", 13, SLATE, False)],
    line_spacing=1.1)
headers = ["Dataset", "A_T", "A_B", "A_N", "HM", "ΔA_N", "ΔHM"]
rows = [
    ["CIFAR-100", "46.42", "61.12", "24.38", "34.85", "+3.6", "+3.0"],
    ["miniImageNet", "37.38", "53.23", "13.60", "21.67", "−2.8", "−3.8"],
    ["CUB-200", "47.98", "64.49", "31.84", "42.63", "+1.0", "−0.0"],
]
cw = [Inches(2.2)] + [Inches(1.65)] * 6
table(s, Inches(0.55), Inches(2.15), Inches(12.1), headers, rows, col_w=cw,
      row_h=Inches(0.52), fs=13, hfs=12, bold_first=True)
bullets(s, Inches(0.55), Inches(4.65), Inches(12), Inches(2.2),
    [
     "**Dataset-dependent, not universal:** helped CIFAR-100 (+3.0 HM), hurt miniImageNet (−3.8 HM), neutral on CUB.",
     "An honest negative-ish result: no single base objective wins everywhere — the right recipe depends on backbone strength.",
     "Reported transparently rather than cherry-picked — a core part of the analysis.",
    ], size=13.5, gap=11, marker="●", marker_col=AMBER)

# ==================================================================
# SLIDE 17 — IMPROVEMENT STUDY (levers)
# ==================================================================
s = slide(); header(s, "Analysis", "Improvement Study — What Actually Works", 17); footer(s)
txt(s, Inches(0.55), Inches(1.35), Inches(12), Inches(0.5),
    [P("Every no-architecture-change lever, measured honestly:", 14, SLATE, False)])
# two columns: works / doesn't
rect(s, Inches(0.55), Inches(1.95), Inches(5.95), Inches(4.6), LIGHT, line=LINE, line_w=Pt(1))
rect(s, Inches(0.55), Inches(1.95), Inches(5.95), Inches(0.5), TEAL)
txt(s, Inches(0.55), Inches(1.95), Inches(5.95), Inches(0.5), [P("✓  Helped", 15, WHITE, True)], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
bullets(s, Inches(0.8), Inches(2.65), Inches(5.5), Inches(3.8),
    [
     "**Transductive rectification** — every metric up, all datasets",
     "**Full-data base prototypes** — helps weak-backbone datasets (mini)",
     "**Best stack: +1.2 to +1.6** overall accuracy, free",
    ], size=13.5, gap=13, marker="▸", marker_col=TEAL)
rect(s, Inches(6.8), Inches(1.95), Inches(5.95), Inches(4.6), LIGHT, line=LINE, line_w=Pt(1))
rect(s, Inches(6.8), Inches(1.95), Inches(5.95), Inches(0.5), RED)
txt(s, Inches(6.8), Inches(1.95), Inches(5.95), Inches(0.5), [P("✗  Did not help", 15, WHITE, True)], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
bullets(s, Inches(7.05), Inches(2.65), Inches(5.5), Inches(3.8),
    [
     "**Novel-logit bias** — only trades base ↔ novel",
     "**Test-time / support augmentation** — no reliable gain",
     "**Feature centering** — negligible effect",
     "**Longer backbone (100 ep)** — blocked by compute (see Slide 19)",
    ], size=13.5, gap=13, marker="▸", marker_col=RED)

# ==================================================================
# SLIDE 18 — KEY FINDING: A_N BOTTLENECK
# ==================================================================
s = slide(); header(s, "Analysis", "Key Finding — Novel-Class Accuracy Is the Bottleneck", 18); footer(s)
s.shapes.add_picture(BASENOVEL, Inches(0.55), Inches(1.5), width=Inches(6.7))
rect(s, Inches(7.55), Inches(1.5), Inches(5.2), Inches(4.9), LIGHT, line=LINE, line_w=Pt(1))
txt(s, Inches(7.75), Inches(1.7), Inches(4.8), Inches(0.4), [P("The central insight", 15, RED, True)])
bullets(s, Inches(7.75), Inches(2.35), Inches(4.8), Inches(4),
    [
     "Across all datasets **A_N (16–31%) ≪ A_B (56–69%)**.",
     "The model remembers base classes well; it **struggles to learn novel ones**.",
     "Overall accuracy and PD hide this — **HM exposes it**.",
     "Every effective lever we found **targets A_N specifically**.",
     "**The ceiling is backbone quality:** a stronger frozen representation lifts A_N and A_B together.",
    ], size=13.5, gap=12, marker="▸", marker_col=RED)

# ==================================================================
# SLIDE 19 — CHALLENGES / LIMITATIONS
# ==================================================================
s = slide(); header(s, "Reflection", "Engineering Challenges, Bugs Fixed & Limitations", 19); footer(s)
rect(s, Inches(0.55), Inches(1.45), Inches(5.95), Inches(2.35), LIGHT, line=LINE, line_w=Pt(1))
txt(s, Inches(0.75), Inches(1.6), Inches(5.5), Inches(0.4), [P("Bugs caught by verifying results", 14, NAVY, True)])
bullets(s, Inches(0.75), Inches(2.15), Inches(5.5), Inches(1.6),
    [
     "**CUB base collapse 1% → 78%** — 100-class joint graph was out-of-distribution for the 15-way GAT; fixed by chunking base prototypes by episode-way.",
     "**Cross-domain crash** — per-sample vs per-class mask shape mismatch; fixed.",
    ], size=12.5, gap=9, marker="▸", marker_col=RED)
rect(s, Inches(6.8), Inches(1.45), Inches(5.95), Inches(2.35), LIGHT, line=LINE, line_w=Pt(1))
txt(s, Inches(7.0), Inches(1.6), Inches(5.5), Inches(0.4), [P("Config-leak fixes", 14, NAVY, True)])
bullets(s, Inches(7.0), Inches(2.15), Inches(5.5), Inches(1.6),
    [
     "Per-dataset hyperparameters leaking across runs (e.g. CUB LR into CIFAR) — fixed with capture-once defaults.",
     "Ablation weights not restoring — fixed with idempotent restore.",
    ], size=12.5, gap=9, marker="▸", marker_col=AMBER)
rect(s, Inches(0.55), Inches(4.0), Inches(12.2), Inches(2.5), LIGHT, line=LINE, line_w=Pt(1))
txt(s, Inches(0.75), Inches(4.15), Inches(11.8), Inches(0.4), [P("Honest limitations", 14, SLATE, True)])
bullets(s, Inches(0.75), Inches(4.7), Inches(11.8), Inches(1.7),
    [
     "Numbers are **single-run, first-pass** on Apple-Silicon — internally consistent and reproducible, but not seed-averaged or tuned to compete with published SOTA.",
     "**Backbone training is compute-capped on this machine** (long background jobs are reaped ~40 min), so the biggest lever — a fully-trained backbone — could not be completed here. It needs a GPU.",
     "The cross-domain protocol is this project's own definition, not a citable standard.",
    ], size=12.5, gap=9, marker="●", marker_col=SLATE)

# ==================================================================
# SLIDE 20 — CONCLUSION & FUTURE WORK
# ==================================================================
s = slide()
rect(s, 0, 0, SW, SH, INK)
rect(s, 0, 0, Inches(0.18), SH, BLUE)
txt(s, Inches(0.8), Inches(0.55), Inches(11), Inches(0.8),
    [P("Conclusion & Future Work", 30, WHITE, True)])
rect(s, Inches(0.85), Inches(1.45), Inches(3.6), Pt(2.5), TEAL)
txt(s, Inches(0.8), Inches(1.7), Inches(11.7), Inches(0.4), [P("WHAT WAS DELIVERED", 13, TEAL, True)])
bullets(s, Inches(0.8), Inches(2.15), Inches(11.7), Inches(2.2),
    [
     "Full STAG-STI pipeline implemented and validated on **3 datasets + cross-domain**, official splits.",
     "Complete **ablation** (every component earns its place) and an honest **improvement study**.",
     "**+1.2 to +1.6** overall accuracy from eval-only levers, no architecture change or retraining.",
     "Central finding: **novel-class accuracy (A_N) is the bottleneck**; the ceiling is backbone quality.",
    ], size=14.5, gap=11, marker="▸", marker_col=WHITE, color=RGBColor(0xE8,0xED,0xF5))
txt(s, Inches(0.8), Inches(4.85), Inches(11.7), Inches(0.4), [P("FUTURE WORK", 13, AMBER, True)])
bullets(s, Inches(0.8), Inches(5.3), Inches(11.7), Inches(1.6),
    [
     "Train the frozen backbone to completion on a **GPU** (100–200 epochs) to lift A_N and A_B together.",
     "Seed-averaged runs and comparison against published FSCIL baselines.",
     "Extend transductive rectification and explore stronger self-supervised base objectives.",
    ], size=14, gap=10, marker="▸", marker_col=AMBER, color=RGBColor(0xE8,0xED,0xF5))
rect(s, 0, SH - Inches(0.22), SW, Inches(0.22), TEAL)

# ---------------- save ----------------
OUT = os.path.join(ROOT, "results", "STAG-STI_Review2.pptx")
prs.save(OUT)
print("Saved", OUT, f"({len(prs.slides.__iter__.__self__._sldIdLst)} slides)")
